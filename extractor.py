# extractor.py
# Dynamically parses the downloaded Excel files from fondbolagen.se and
# transforms data into the SIFA master CSV format.
#
# The extractor is context-aware: it dynamically identifies section headers,
# category rows, and quarter columns without hardcoding row numbers.
# This allows it to adapt if the source structure changes.
#
# Source Excel layout (per fund type section):
#   Section header row:    "All types of funds" / "Equity funds" / etc.
#   Column header row:     Quarter 1 | Quarter 2 | Quarter 3 | Quarter 4 | Net savings | Net savings | Net assets | Net assets
#   Sub-header row:        (empty)    | (empty)   | (empty)   | (empty)   | sum         | %           | <date>     | %
#   Data rows (10):        Category   | Q1 val    | Q2 val    | Q3 val    | Q4 val      | Sum         | %          | Assets     | Assets %
#   ... (9 categories + TOTAL)
#
# Values in cells may be formulas — we use data_only=True to read computed values.

import os
import re
import logging
from collections import OrderedDict

import openpyxl
from openpyxl.utils import get_column_letter
import pandas as pd

import config

logger = logging.getLogger(__name__)


# ─────────────────────────────────────────────────────────────────────────────
# Column visibility
# ─────────────────────────────────────────────────────────────────────────────

def _get_visible_cols(ws):
    """
    Return the set of 1-based column indices that are NOT hidden.
    Columns with no dimension entry are treated as visible (Excel default).
    Only active when config.SKIP_HIDDEN_COLUMNS is True.
    """
    visible = set()
    for idx in range(1, ws.max_column + 1):
        letter = get_column_letter(idx)
        dim = ws.column_dimensions.get(letter)
        if dim is None or not dim.hidden:
            visible.add(idx)
    return visible


# ─────────────────────────────────────────────────────────────────────────────
# Dynamic section finder
# ─────────────────────────────────────────────────────────────────────────────

def _find_sections(ws):
    """
    Dynamically locate fund type sections in the worksheet.
    Searches column A for known section header names.

    Returns list of dicts:
        [{'fund_type': 'ALLTYPES', 'header_row': 8, 'data_start_row': 11}, ...]
    """
    sections = []
    fund_type_lookup = {}
    for ft_code, ft_name in config.FUND_TYPE_DISPLAY.items():
        fund_type_lookup[ft_name.lower()] = ft_code

    for row in range(1, ws.max_row + 1):
        cell_val = ws.cell(row=row, column=1).value
        if cell_val is None:
            continue

        cell_text = str(cell_val).strip().lower()

        for ft_name_lower, ft_code in fund_type_lookup.items():
            if cell_text == ft_name_lower:
                # Found a section header — now find where data rows start
                # Data rows start after the header row and sub-header row
                # Look for the first row below that has a category label in column A
                data_start = None
                for r in range(row + 1, min(row + 10, ws.max_row + 1)):
                    v = ws.cell(row=r, column=1).value
                    if v is not None and str(v).strip() and str(v).strip() != 'TOTAL':
                        # Verify it's a category name, not a column header
                        if 'Quarter' not in str(v) and 'Net' not in str(v):
                            data_start = r
                            break

                if data_start is not None:
                    sections.append({
                        'fund_type': ft_code,
                        'header_row': row,
                        'data_start_row': data_start,
                    })
                    logger.debug(f'Section "{ft_code}" at row {row}, '
                                 f'data starts at row {data_start}')
                break

    if len(sections) != len(config.FUND_TYPES):
        found = [s['fund_type'] for s in sections]
        expected = config.FUND_TYPES
        logger.warning(f'Expected {len(expected)} sections, found {len(found)}: '
                       f'{found}')

    return sections


def _find_category_rows(ws, data_start_row):
    """
    Starting from data_start_row, read 10 consecutive rows (9 categories + TOTAL).
    Returns list of row numbers in order.
    """
    rows = []
    row = data_start_row
    while len(rows) < 10 and row <= ws.max_row:
        cell_val = ws.cell(row=row, column=1).value
        if cell_val is not None:
            label = str(cell_val).strip()
            if label:
                rows.append(row)
                if label == 'TOTAL':
                    break
        row += 1

    if len(rows) != 10:
        logger.warning(f'Expected 10 category rows from row {data_start_row}, '
                       f'found {len(rows)}')

    return rows


# ─────────────────────────────────────────────────────────────────────────────
# Quarter detection
# ─────────────────────────────────────────────────────────────────────────────

def _detect_quarters_with_data(ws, data_rows, visible_cols=None):
    """
    Check which quarters (columns B-E) have actual data.
    Examines all non-TOTAL data rows to see which quarter columns
    contain at least one non-zero numeric value.

    When visible_cols is provided (SKIP_HIDDEN_COLUMNS=True), any quarter
    whose source column is hidden is excluded regardless of its values —
    the provider hides columns to signal data is not yet published.

    Returns sorted list of quarter numbers [1, 2, 3] or [1, 2, 3, 4].
    """
    if not data_rows:
        return []

    # Check all non-TOTAL rows (exclude last row which is TOTAL)
    check_rows = data_rows[:-1] if len(data_rows) > 1 else data_rows

    quarters = []
    for q_num, col_idx in config.EXCEL_QUARTER_COLS.items():
        if visible_cols is not None and col_idx not in visible_cols:
            logger.debug(f'Q{q_num} (col {col_idx}) is hidden — skipping')
            continue
        has_nonzero = False
        for row_num in check_rows:
            val = ws.cell(row=row_num, column=col_idx).value
            if val is not None and isinstance(val, (int, float)) and val != 0:
                has_nonzero = True
                break
        if has_nonzero:
            quarters.append(q_num)

    return sorted(quarters)


# ─────────────────────────────────────────────────────────────────────────────
# Value extraction
# ─────────────────────────────────────────────────────────────────────────────

def _copy_values_sheet(wb):
    """
    Copy all cell values from the active sheet to a new sheet.

    This is the 'copy method': loading with data_only=True gives us
    cached formula results, and copying them to a fresh sheet ensures
    we get clean, raw numeric values — no formulas, no artifacts.
    """
    source_ws = wb.active
    target_title = f'{source_ws.title}_Values'
    target_ws = wb.create_sheet(title=target_title)

    for row in source_ws.iter_rows():
        for cell in row:
            target_ws.cell(row=cell.row, column=cell.column, value=cell.value)

    return target_ws


def _get_cell_value(ws, row, col):
    """
    Read a cell value, returning None for empty/non-numeric cells.
    Formats to 15 significant digits (:.15g) to match Excel's formula bar
    precision and strip IEEE 754 tail noise.
    """
    val = ws.cell(row=row, column=col).value
    if val is None:
        return None
    if isinstance(val, (int, float)):
        return float(f'{float(val):.15g}')
    # Try to parse string representation
    try:
        cleaned = str(val).replace(',', '').strip()
        if cleaned == '' or cleaned == '-':
            return None
        return float(f'{float(cleaned):.15g}')
    except (ValueError, TypeError):
        return None


def _extract_section_data(ws, section, quarter_num, is_last_quarter, visible_cols=None):
    """
    Extract data for one fund type section for one quarter.

    Returns dict mapping metric positions to values.
    The dict has keys 0-49 representing the 50 column positions
    within this fund type's block.

    For quarterly net savings (positions 0-9): always populated.
    For sum/pct/assets (positions 10-49): only populated if is_last_quarter
    AND the source column is visible (when SKIP_HIDDEN_COLUMNS is True).
    """
    data_rows = _find_category_rows(ws, section['data_start_row'])

    if len(data_rows) < 10:
        logger.warning(f'Section {section["fund_type"]}: only {len(data_rows)} '
                       f'category rows found')
        while len(data_rows) < 10:
            data_rows.append(None)

    quarter_col = config.EXCEL_QUARTER_COLS[quarter_num]
    values = {}

    # Positions 0-9: Net savings for this quarter
    for i, row_num in enumerate(data_rows):
        if row_num is not None:
            values[i] = _get_cell_value(ws, row_num, quarter_col)
        else:
            values[i] = None

    if is_last_quarter:
        summary_cols = [
            (10, config.EXCEL_SUM_COL,       'NETSAVINGSUM  (F)'),
            (20, config.EXCEL_PCT_COL,        'NETSAVINGPERC (G)'),
            (30, config.EXCEL_ASSET_COL,      'NETASSET      (H)'),
            (40, config.EXCEL_ASSET_PCT_COL,  'NETASSETPERC  (I)'),
        ]
        for base_pos, col_idx, label in summary_cols:
            if visible_cols is not None and col_idx not in visible_cols:
                logger.debug(f'Section {section["fund_type"]}: {label} is hidden — skipping')
                continue
            for i, row_num in enumerate(data_rows):
                if row_num is not None:
                    values[base_pos + i] = _get_cell_value(ws, row_num, col_idx)
                else:
                    values[base_pos + i] = None

    return values


# ─────────────────────────────────────────────────────────────────────────────
# Main extraction
# ─────────────────────────────────────────────────────────────────────────────

def extract(excel_path, year):
    """
    Extract quarterly data from a source Excel file.

    Args:
        excel_path: Path to the downloaded Excel file
        year: The year this file represents (int)

    Returns:
        dict mapping quarter labels ('YYYY-QN') to row data dicts.
        Each row data dict maps column SIMBA codes to values.
    """
    logger.info(f'Extracting data from: {excel_path} (year {year})')

    wb = openpyxl.load_workbook(excel_path, data_only=True)

    # Read column visibility from the original sheet BEFORE copying —
    # the copy method transfers cell values only, not column dimensions.
    source_ws = wb.active
    visible_cols_raw = _get_visible_cols(source_ws) if config.SKIP_HIDDEN_COLUMNS else None

    # Copy method: copy all values to a fresh sheet to get clean raw values
    ws = _copy_values_sheet(wb)
    logger.info(f'Sheet: {source_ws.title} -> {ws.title}, '
                f'{ws.max_row} rows x {ws.max_column} cols')

    # Step 1: Dynamically find all fund type sections
    sections = _find_sections(ws)
    if not sections:
        logger.error('No fund type sections found in the Excel file')
        return {}

    # Ensure sections are in the expected order
    section_order = {ft: i for i, ft in enumerate(config.FUND_TYPES)}
    sections.sort(key=lambda s: section_order.get(s['fund_type'], 999))

    logger.info(f'Found {len(sections)} sections: '
                f'{[s["fund_type"] for s in sections]}')

    # Step 2: Build column visibility map (respects provider publication signal)
    visible_cols = visible_cols_raw
    if visible_cols is not None:
        hidden = sorted(set(range(1, ws.max_column + 1)) - visible_cols)
        logger.info(f'SKIP_HIDDEN_COLUMNS=True | visible cols: {sorted(visible_cols)} '
                    f'| hidden cols: {hidden}')

    # Step 3: Detect which quarters have data
    first_section = sections[0]
    first_data_rows = _find_category_rows(ws, first_section['data_start_row'])
    quarters = _detect_quarters_with_data(ws, first_data_rows, visible_cols)

    if not quarters:
        logger.warning('No quarters with data found')
        return {}

    logger.info(f'Quarters with data: {quarters}')
    last_quarter = max(quarters)

    # Step 4: Extract data for each quarter
    result = OrderedDict()

    for q_num in quarters:
        quarter_label = f'{year}-Q{q_num}'
        is_last = (q_num == last_quarter)

        row_values = {}
        col_offset = 0

        for section in sections:
            section_data = _extract_section_data(ws, section, q_num, is_last, visible_cols)

            # Map section positions (0-49) to global DATA_COLUMNS indices
            for pos, value in section_data.items():
                global_idx = col_offset + pos
                if global_idx < len(config.DATA_COLUMNS):
                    col_code = config.DATA_COLUMNS[global_idx]
                    row_values[col_code] = value

            col_offset += 50  # 50 columns per fund type section

        # Count populated values
        populated = sum(1 for v in row_values.values() if v is not None)
        logger.info(f'{quarter_label}: {populated}/{len(config.DATA_COLUMNS)} '
                    f'columns populated'
                    f'{" (last quarter - includes sum/assets)" if is_last else ""}')

        result[quarter_label] = row_values

    wb.close()
    return result


def extract_multiple(file_list):
    """
    Extract data from multiple year files.

    Args:
        file_list: list of (year, filepath) tuples

    Returns:
        Combined dict mapping quarter labels to row data dicts.
    """
    all_data = OrderedDict()

    for year, filepath in sorted(file_list, key=lambda x: x[0]):
        try:
            year_data = extract(filepath, year)
            all_data.update(year_data)
        except Exception as e:
            logger.error(f'Failed to extract {year} from {filepath}: {e}')

    logger.info(f'Total quarters extracted: {len(all_data)}')
    return all_data
