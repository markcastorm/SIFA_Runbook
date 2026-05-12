"""
SIFA Comparison Tool
=====================
Dynamic path discovery - no manual path entry needed.

What it auto-discovers:
  - Pipeline output:    Latest run in ../output/latest/  (SIFA_DATA.csv)
  - Source downloads:   Downloaded Excel files in ../downloads/  (for verification)
  - Reference file:     Any .xlsx/.csv dropped into ./reference_input/

Four separate CSV reports:
    report_mismatches_full.csv     -- large value conflicts (wrong mapping / value)
    report_mismatches_decimal.csv  -- small rounding differences
    report_missing.csv             -- in reference but absent from pipeline
    report_new_values.csv          -- in pipeline but absent from reference

Annotated Excel (if xlsx inputs):
    Red    = full mismatch
    Orange = decimal mismatch
    Yellow = missing
    Green  = new value

Usage:
  1. Drop your reference .xlsx/.csv into the  reference_input/  folder
  2. Run: python compare.py
  3. Find results in the new run_YYYYMMDD_HHMMSS/ folder

Settings (tolerances, verbosity) are in config_compare.py.
"""

import openpyxl
from openpyxl.styles import PatternFill, Font
import csv
import os
import sys
import glob
import re
from datetime import datetime
from collections import defaultdict

import config_compare as config


# ---------------------------------------------------------------------------
# Path Discovery
# ---------------------------------------------------------------------------

def find_pipeline_output(project_root):
    """
    Find the latest pipeline DATA file in output/latest/.

    Prefers CSV over xlsx because:
      - CSV contains ALL quarters (full master data)
      - xlsx only contains new quarters from the latest run
    """
    latest_dir = os.path.join(project_root, 'output', 'latest')

    if not os.path.isdir(latest_dir):
        raise FileNotFoundError(
            f'No output/latest/ directory found: {latest_dir}\n'
            'Run the pipeline first to generate output.'
        )

    # Prefer CSV (has full master data with all quarters)
    csv_files = glob.glob(os.path.join(latest_dir, 'SIFA_DATA.csv'))
    if csv_files:
        return csv_files[0]

    # Fall back to xlsx (only has new quarters from latest run)
    xlsx_files = glob.glob(os.path.join(latest_dir, 'SIFA_DATA_latest.xlsx'))
    if not xlsx_files:
        xlsx_files = glob.glob(os.path.join(latest_dir, 'SIFA_DATA_*.xlsx'))

    if not xlsx_files:
        raise FileNotFoundError(
            f'No SIFA_DATA.csv or SIFA_DATA_*.xlsx found in: {latest_dir}'
        )

    return xlsx_files[0]


def find_source_downloads(project_root):
    """
    Find downloaded source Excel files in ../downloads/.
    Returns dict: {year: filepath, ...}
    """
    download_root = os.path.join(project_root, 'downloads')
    if not os.path.isdir(download_root):
        return {}

    source_map = {}

    for root, dirs, files in os.walk(download_root):
        for f in files:
            if f.endswith(('.xlsx', '.xls')) and not f.startswith('~$'):
                year_match = re.search(r'\b(20\d{2})\b', os.path.join(root, f))
                if year_match:
                    year = int(year_match.group(1))
                    filepath = os.path.join(root, f)
                    if year not in source_map or os.path.getmtime(filepath) > os.path.getmtime(source_map[year]):
                        source_map[year] = filepath

    return source_map


def find_reference_file(compare_dir):
    """Find the reference file in reference_input/."""
    ref_dir = os.path.join(compare_dir, 'reference_input')
    os.makedirs(ref_dir, exist_ok=True)

    ref_files = (
        glob.glob(os.path.join(ref_dir, '*.xlsx')) +
        glob.glob(os.path.join(ref_dir, '*.csv'))
    )
    ref_files = [f for f in ref_files if not os.path.basename(f).startswith('~$')]

    if not ref_files:
        raise FileNotFoundError(
            f'No reference .xlsx or .csv found in:\n  {ref_dir}\n\n'
            'Drop your reference file there and re-run.'
        )

    if len(ref_files) > 1:
        names = '\n  '.join(os.path.basename(f) for f in ref_files)
        raise ValueError(
            f'Multiple reference files found in reference_input/ - keep only one:\n  {names}'
        )

    return ref_files[0]


# ---------------------------------------------------------------------------
# Data Loaders
# ---------------------------------------------------------------------------

def _load_sifa_xlsx(path, label):
    """
    Load a SIFA DATA Excel file.

    Layout:
        Row 1: SIMBA codes (col A empty, B+ = codes)
        Row 2: descriptions
        Row 3+: data rows (col A = quarter label like '2024-Q1')

    Returns: (codes_list, cells_dict, descriptions_list)
    """
    if config.VERBOSE:
        print(f'Loading {label}: {os.path.basename(path)}')

    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb.active

    codes = []
    for c in range(2, ws.max_column + 1):
        code = ws.cell(row=1, column=c).value
        if code is None:
            break
        codes.append(str(code).strip())

    descriptions = []
    for c in range(2, 2 + len(codes)):
        desc = ws.cell(row=2, column=c).value
        descriptions.append(str(desc).strip() if desc else '')

    cells = {}
    quarters_found = []
    for r in range(3, ws.max_row + 1):
        quarter_val = ws.cell(row=r, column=1).value
        if quarter_val is None:
            continue
        ql = str(quarter_val).strip()
        if not ql:
            continue
        quarters_found.append(ql)

        for c_idx in range(len(codes)):
            v = ws.cell(row=r, column=c_idx + 2).value
            if v is not None:
                cells[(ql, c_idx)] = v

    wb.close()

    if config.VERBOSE:
        print(f'  Loaded: {len(codes)} codes, {len(cells)} cells, '
              f'{len(quarters_found)} quarters')

    return codes, cells, descriptions


def _load_sifa_csv(path, label):
    """
    Load a SIFA DATA CSV file (with description sub-header row).

    Layout:
        Row 1: empty + SIMBA codes
        Row 2: empty + descriptions
        Row 3+: quarter_label + values

    Returns: (codes_list, cells_dict, descriptions_list)
    """
    if config.VERBOSE:
        print(f'Loading {label}: {os.path.basename(path)}')

    with open(path, 'r', newline='', encoding='utf-8') as f:
        reader = csv.reader(f)
        header = next(reader, None)
        second = next(reader, None)

        codes = [c.strip() for c in header[1:]] if header else []

        descriptions = []
        data_rows_raw = []
        if second and len(second) > 1 and ':' in str(second[1]):
            descriptions = [d.strip() for d in second[1:]]
            for row in reader:
                if row and row[0].strip():
                    data_rows_raw.append(row)
        else:
            if second and second[0].strip():
                data_rows_raw.append(second)
            for row in reader:
                if row and row[0].strip():
                    data_rows_raw.append(row)

    cells = {}
    quarters_found = []
    for row in data_rows_raw:
        ql = row[0].strip()
        quarters_found.append(ql)
        for c_idx, val in enumerate(row[1:]):
            if c_idx < len(codes) and val.strip() != '':
                try:
                    cells[(ql, c_idx)] = float(val)
                except ValueError:
                    cells[(ql, c_idx)] = val.strip()

    if config.VERBOSE:
        print(f'  Loaded: {len(codes)} codes, {len(cells)} cells, '
              f'{len(quarters_found)} quarters')

    return codes, cells, descriptions


def load_data_file(path, label):
    """Load a SIFA data file (xlsx or csv)."""
    if path.lower().endswith('.csv'):
        return _load_sifa_csv(path, label)
    else:
        return _load_sifa_xlsx(path, label)


# ---------------------------------------------------------------------------
# Source Excel Verification
# ---------------------------------------------------------------------------

def _load_source_sections(ws):
    """
    Locate fund type sections in a source Excel worksheet.
    Returns list of {'fund_type_display': name, 'header_row': N, 'data_start': N}
    """
    section_names = [
        'all types of funds',
        'equity funds',
        'balanced funds',
        'long term fixed income funds',
        'short term fixed income funds',
        'hedge funds',
        'other funds',
    ]

    sections = []
    for row in range(1, ws.max_row + 1):
        cell_val = ws.cell(row=row, column=1).value
        if cell_val is None:
            continue
        cell_text = str(cell_val).strip().lower()

        for sn in section_names:
            if cell_text == sn:
                data_start = None
                for r in range(row + 1, min(row + 10, ws.max_row + 1)):
                    v = ws.cell(row=r, column=1).value
                    if v is not None and str(v).strip():
                        if 'Quarter' not in str(v) and 'Net' not in str(v):
                            data_start = r
                            break
                if data_start is not None:
                    sections.append({
                        'fund_type_display': sn,
                        'header_row': row,
                        'data_start': data_start,
                    })
                break

    return sections


def verify_from_source(source_path, quarter_label, col_idx, pipeline_val, ref_val):
    """
    Open the source Excel file and check the actual value for a specific
    quarter/column to determine which side (pipeline or reference) is correct.
    """
    match = re.match(r'(\d{4})-Q(\d)', quarter_label)
    if not match:
        return {'source_value': None, 'note': f'Cannot parse quarter: {quarter_label}'}

    year = int(match.group(1))
    q_num = int(match.group(2))

    try:
        wb = openpyxl.load_workbook(source_path, data_only=True)
        ws = wb.active
    except Exception as e:
        return {'source_value': None, 'note': f'Cannot open source: {e}'}

    section_idx = col_idx // 50
    position_in_section = col_idx % 50
    metric_idx = position_in_section // 10
    category_idx = position_in_section % 10

    sections = _load_source_sections(ws)
    if section_idx >= len(sections):
        wb.close()
        return {'source_value': None, 'note': f'Section {section_idx} not found in source'}

    section = sections[section_idx]

    data_rows = []
    row = section['data_start']
    while len(data_rows) < 10 and row <= ws.max_row:
        cell_val = ws.cell(row=row, column=1).value
        if cell_val is not None and str(cell_val).strip():
            data_rows.append(row)
            if str(cell_val).strip() == 'TOTAL':
                break
        row += 1

    if category_idx >= len(data_rows):
        wb.close()
        return {'source_value': None, 'note': f'Category row {category_idx} not found'}

    target_row = data_rows[category_idx]

    quarter_col_map = {1: 2, 2: 3, 3: 4, 4: 5}

    if metric_idx == 0:
        target_col = quarter_col_map.get(q_num, 2)
    elif metric_idx == 1:
        target_col = 6
    elif metric_idx == 2:
        target_col = 7
    elif metric_idx == 3:
        target_col = 8
    elif metric_idx == 4:
        target_col = 9
    else:
        wb.close()
        return {'source_value': None, 'note': f'Unknown metric index: {metric_idx}'}

    source_value = ws.cell(row=target_row, column=target_col).value
    wb.close()

    if isinstance(source_value, (int, float)):
        source_value = float(source_value)
    elif source_value is not None:
        try:
            source_value = float(str(source_value).replace(',', '').strip())
        except (ValueError, TypeError):
            pass

    result = {
        'source_value': source_value,
        'source_row': target_row,
        'source_col': target_col,
        'section': section['fund_type_display'],
        'pipeline_correct': None,
        'reference_correct': None,
        'note': '',
    }

    if source_value is None:
        result['note'] = 'Source cell is empty'
        return result

    def _values_match(a, b):
        if a is None or b is None:
            return a is None and b is None
        try:
            fa, fb = float(a), float(b)
            return abs(fa - fb) < config.FLOAT_TOLERANCE
        except (ValueError, TypeError):
            return str(a) == str(b)

    p_match = _values_match(source_value, pipeline_val)
    r_match = _values_match(source_value, ref_val)

    result['pipeline_correct'] = p_match
    result['reference_correct'] = r_match

    if p_match and r_match:
        result['note'] = 'Both match source (within tolerance)'
    elif p_match and not r_match:
        result['note'] = 'Pipeline matches source; reference differs'
    elif not p_match and r_match:
        result['note'] = 'Reference matches source; pipeline differs'
    else:
        result['note'] = 'Neither matches source exactly'

    return result


# ---------------------------------------------------------------------------
# Comparison Engine
# ---------------------------------------------------------------------------

class ComparisonReport:

    def __init__(self, pipeline_path, reference_path, source_downloads):
        self.pipeline_path  = pipeline_path
        self.reference_path = reference_path
        self.source_downloads = source_downloads  # {year: filepath}

        self.pipeline_codes  = []
        self.pipeline_cells  = {}
        self.pipeline_descs  = []
        self.reference_codes = []
        self.reference_cells = {}
        self.reference_descs = []

        self.matches            = []
        self.mismatches_full    = []
        self.mismatches_decimal = []
        self.missing            = []
        self.extra              = []

        self.full_mismatch_coords    = []
        self.decimal_mismatch_coords = []
        self.missing_coords          = []
        self.extra_coords            = []

    # ------------------------------------------------------------------
    # Helpers
    # ------------------------------------------------------------------

    def get_section_name(self, col_idx):
        section_idx = col_idx // 50
        section_names = [
            'All types of funds', 'Equity funds', 'Balanced funds',
            'Long term fixed income', 'Short term fixed income',
            'Hedge funds', 'Other funds',
        ]
        if section_idx < len(section_names):
            return section_names[section_idx]
        return 'UNKNOWN'

    def get_metric_name(self, col_idx):
        position = col_idx % 50
        metric_idx = position // 10
        metric_names = [
            'Net savings', 'Net savings sum', 'Net savings %',
            'Net assets', 'Net assets %',
        ]
        if metric_idx < len(metric_names):
            return metric_names[metric_idx]
        return 'UNKNOWN'

    def get_category_name(self, col_idx):
        position = col_idx % 50
        cat_idx = position % 10
        cat_names = [
            'Swedish households', 'ISK', 'IPS', 'Unit linked',
            'Premium Pension', 'Nominee Accounts',
            'Non profit inst.', 'Swedish corporations',
            'Others', 'TOTAL',
        ]
        if cat_idx < len(cat_names):
            return cat_names[cat_idx]
        return 'UNKNOWN'

    def compare_values(self, val1, val2):
        if val1 is None and val2 is None:
            return True, 0
        if val1 is None or val2 is None:
            return False, None
        try:
            f1, f2 = float(val1), float(val2)
            diff = f1 - f2
            if abs(diff) < config.FLOAT_TOLERANCE:
                return True, 0
            return False, diff
        except (ValueError, TypeError):
            return str(val1) == str(val2), None

    def build_row_dict(self, status, quarter, col_idx, code, out_val, ref_val, diff):
        section  = self.get_section_name(col_idx)
        metric   = self.get_metric_name(col_idx)
        category = self.get_category_name(col_idx)
        desc = ''
        if col_idx < len(self.pipeline_descs):
            desc = self.pipeline_descs[col_idx]

        year_match = re.match(r'(\d{4})', quarter)
        year = int(year_match.group(1)) if year_match else None
        source_file = ''
        if year and year in self.source_downloads:
            source_file = os.path.basename(self.source_downloads[year])

        return {
            'Pipeline_File':     os.path.basename(self.pipeline_path),
            'Reference_File':    os.path.basename(self.reference_path),
            'Status':            status,
            'Quarter':           quarter,
            'Excel_Column':      col_idx + 2,
            'Column_Code':       code or '',
            'Section':           section,
            'Category':          category,
            'Metric':            metric,
            'Column_Description': desc,
            'Pipeline_Value':    out_val if out_val is not None else '',
            'Reference_Value':   ref_val if ref_val is not None else '',
            'Difference':        round(diff, 6) if diff is not None else '',
            'Source_File':       source_file if source_file else 'NO SOURCE FILE',
            '_year':             year,
            '_col_idx':          col_idx,
        }

    # ------------------------------------------------------------------
    # Core comparison
    # ------------------------------------------------------------------

    def perform_comparison(self):
        if config.VERBOSE:
            print('\nPerforming comparison...')

        pipeline_quarters = set(k[0] for k in self.pipeline_cells.keys())
        reference_quarters = set(k[0] for k in self.reference_cells.keys())
        common_quarters = sorted(pipeline_quarters & reference_quarters)

        if config.VERBOSE:
            print(f'  Pipeline quarters:  {sorted(pipeline_quarters)}')
            print(f'  Reference quarters: {sorted(reference_quarters)}')
            print(f'  Overlapping:        {common_quarters}')

        pipeline_only  = sorted(pipeline_quarters - reference_quarters)
        reference_only = sorted(reference_quarters - pipeline_quarters)

        if pipeline_only and config.VERBOSE:
            print(f'  Pipeline only (not compared): {pipeline_only}')
        if reference_only and config.VERBOSE:
            print(f'  Reference only (not compared): {reference_only}')

        num_codes = max(len(self.pipeline_codes), len(self.reference_codes))

        for quarter in common_quarters:
            for col_idx in range(num_codes):
                out_val = self.pipeline_cells.get((quarter, col_idx))
                ref_val = self.reference_cells.get((quarter, col_idx))
                code = self.pipeline_codes[col_idx] if col_idx < len(self.pipeline_codes) else ''

                if out_val is None and ref_val is None:
                    continue
                elif out_val is not None and ref_val is not None:
                    match, diff = self.compare_values(out_val, ref_val)
                    if match:
                        self.matches.append((quarter, col_idx))
                    else:
                        row_dict = self.build_row_dict(
                            'MISMATCH', quarter, col_idx,
                            code, out_val, ref_val, diff)

                        # Separate into full vs decimal mismatch
                        abs_diff = abs(float(diff)) if diff is not None else 0
                        if abs_diff > config.DECIMAL_MISMATCH_THRESHOLD:
                            row_dict['Status'] = 'MISMATCH_FULL'
                            self.mismatches_full.append(row_dict)
                            self.full_mismatch_coords.append((quarter, col_idx))
                        else:
                            row_dict['Status'] = 'MISMATCH_DECIMAL'
                            self.mismatches_decimal.append(row_dict)
                            self.decimal_mismatch_coords.append((quarter, col_idx))

                elif out_val is None:
                    self.missing.append(
                        self.build_row_dict('MISSING', quarter, col_idx,
                                            code, None, ref_val, None)
                    )
                    self.missing_coords.append((quarter, col_idx))
                else:
                    self.extra.append(
                        self.build_row_dict('NEW_VALUE', quarter, col_idx,
                                            code, out_val, None, None)
                    )
                    self.extra_coords.append((quarter, col_idx))

        if config.VERBOSE:
            total_mm = len(self.mismatches_full) + len(self.mismatches_decimal)
            print(f'\n  Matches:             {len(self.matches)}')
            print(f'  Mismatches (total):  {total_mm}')
            print(f'    Full mismatches:   {len(self.mismatches_full)}')
            print(f'    Decimal mismatches:{len(self.mismatches_decimal)}')
            print(f'  Missing:             {len(self.missing)}')
            print(f'  New values:          {len(self.extra)}')

    # ------------------------------------------------------------------
    # Source verification for mismatches
    # ------------------------------------------------------------------

    def verify_mismatches(self):
        """Check each mismatch against the source Excel to determine who is correct."""
        all_mismatches = self.mismatches_full + self.mismatches_decimal
        if not all_mismatches or not self.source_downloads:
            return

        if config.VERBOSE:
            print('\nVerifying mismatches against source Excel files...')

        verified = 0
        for row in all_mismatches:
            year = row.get('_year')
            if year and year in self.source_downloads:
                result = verify_from_source(
                    self.source_downloads[year],
                    row['Quarter'],
                    row['_col_idx'],
                    row['Pipeline_Value'],
                    row['Reference_Value'],
                )
                row['Source_Value'] = result.get('source_value', '')
                row['Pipeline_Correct'] = result.get('pipeline_correct', '')
                row['Reference_Correct'] = result.get('reference_correct', '')
                row['Verification_Note'] = result.get('note', '')
                verified += 1
            else:
                row['Source_Value'] = ''
                row['Pipeline_Correct'] = ''
                row['Reference_Correct'] = ''
                row['Verification_Note'] = 'No source file available'

        if config.VERBOSE:
            print(f'  Verified {verified}/{len(all_mismatches)} mismatches')

            p_correct = sum(1 for r in all_mismatches if r.get('Pipeline_Correct') is True)
            r_correct = sum(1 for r in all_mismatches if r.get('Reference_Correct') is True)
            both = sum(1 for r in all_mismatches
                       if r.get('Pipeline_Correct') is True and r.get('Reference_Correct') is True)
            neither = sum(1 for r in all_mismatches
                          if r.get('Pipeline_Correct') is False and r.get('Reference_Correct') is False)

            print(f'  Pipeline correct:  {p_correct}')
            print(f'  Reference correct: {r_correct}')
            if both:
                print(f'  Both correct (tolerance): {both}')
            if neither:
                print(f'  Neither correct: {neither}')

    # ------------------------------------------------------------------
    # CSV writers
    # ------------------------------------------------------------------

    _CSV_HEADER = [
        'Pipeline_File', 'Reference_File',
        'Status', 'Quarter', 'Excel_Column', 'Column_Code',
        'Section', 'Category', 'Metric',
        'Column_Description',
        'Pipeline_Value', 'Reference_Value', 'Difference',
        'Source_File',
        'Source_Value', 'Pipeline_Correct', 'Reference_Correct',
        'Verification_Note',
    ]

    def _write_category_csv(self, rows, output_path, label):
        if config.VERBOSE:
            print(f'  Writing {label} report: {os.path.basename(output_path)}  '
                  f'({len(rows)} rows)')

        with open(output_path, 'w', newline='', encoding='utf-8') as f:
            writer = csv.DictWriter(f, fieldnames=self._CSV_HEADER,
                                    extrasaction='ignore')
            writer.writeheader()
            for row in rows:
                writer.writerow(row)

    # ------------------------------------------------------------------
    # Summary report
    # ------------------------------------------------------------------

    def generate_summary_stats(self):
        total_mm = len(self.mismatches_full) + len(self.mismatches_decimal)
        total = len(self.matches) + total_mm + len(self.missing) + len(self.extra)
        return {
            'timestamp':            datetime.now().strftime('%Y-%m-%d %H:%M:%S'),
            'pipeline_file':        self.pipeline_path,
            'reference_file':       self.reference_path,
            'total_cells_compared': total,
            'matches':              len(self.matches),
            'mismatches_full':      len(self.mismatches_full),
            'mismatches_decimal':   len(self.mismatches_decimal),
            'mismatches_total':     total_mm,
            'missing':              len(self.missing),
            'new_values':           len(self.extra),
            'match_rate':           (len(self.matches) / total * 100) if total > 0 else 0,
        }

    def analyze_by_category(self, rows, key_func):
        groups = defaultdict(list)
        for row in rows:
            groups[key_func(row)].append(row)
        return groups

    def _write_mismatch_detail(self, f, label, mismatches, color_label):
        """Write mismatch detail section to summary report."""
        if not mismatches:
            return

        f.write(f'\n  {label} ({color_label} in annotated Excel):\n')
        f.write(f'  Count: {len(mismatches)}\n')

        f.write('\n    By Section:\n')
        by_section = self.analyze_by_category(
            mismatches, lambda r: r['Section'])
        for section, items in sorted(by_section.items()):
            f.write(f"      {section}: {len(items)}\n")

        f.write('\n    By Metric:\n')
        by_metric = self.analyze_by_category(
            mismatches, lambda r: r['Metric'])
        for metric, items in sorted(by_metric.items(),
                                    key=lambda x: -len(x[1])):
            f.write(f"      {metric}: {len(items)}\n")

        f.write('\n    By Quarter:\n')
        by_quarter = self.analyze_by_category(
            mismatches, lambda r: r['Quarter'])
        for quarter, items in sorted(by_quarter.items()):
            f.write(f"      {quarter}: {len(items)}\n")

        # Source verification summary
        p_correct = sum(1 for r in mismatches
                        if r.get('Pipeline_Correct') is True)
        r_correct = sum(1 for r in mismatches
                        if r.get('Reference_Correct') is True)
        if p_correct or r_correct:
            f.write('\n    Source Verification:\n')
            f.write(f"      Pipeline matches source: {p_correct}\n")
            f.write(f"      Reference matches source: {r_correct}\n")

        f.write(f'\n    Top examples (by absolute difference):\n')
        sorted_mm = sorted(
            [r for r in mismatches if r['Difference'] != ''],
            key=lambda r: abs(float(r['Difference']))
            if r['Difference'] else 0,
            reverse=True,
        )[:config.MAX_CONSOLE_EXAMPLES]
        for r in sorted_mm:
            f.write(f"      Quarter: {r['Quarter']}  Col: {r['Excel_Column']}  "
                    f"{r['Section']} | {r['Category']} | {r['Metric']}\n")
            f.write(f"        Pipeline : {r['Pipeline_Value']}\n")
            f.write(f"        Reference: {r['Reference_Value']}\n")
            f.write(f"        Diff     : {r['Difference']}\n")
            if r.get('Source_Value') != '':
                f.write(f"        Source   : {r.get('Source_Value', '')}\n")
                f.write(f"        Verdict  : {r.get('Verification_Note', '')}\n")
        f.write('\n')

    def write_summary_report(self, output_path):
        if config.VERBOSE:
            print(f'Writing summary report: {os.path.basename(output_path)}')

        stats = self.generate_summary_stats()

        with open(output_path, 'w', encoding='utf-8') as f:
            f.write('=' * 70 + '\n')
            f.write('SIFA COMPARISON REPORT\n')
            f.write('=' * 70 + '\n\n')
            f.write(f"Generated:  {stats['timestamp']}\n\n")

            f.write('FILES COMPARED:\n')
            f.write(f"  Pipeline Output : {os.path.basename(stats['pipeline_file'])}\n")
            f.write(f"  Reference       : {os.path.basename(stats['reference_file'])}\n")
            f.write(f"  Source files     : {len(self.source_downloads)} year(s)\n\n")

            f.write('OVERALL STATISTICS:\n')
            f.write(f"  Total cells compared : {stats['total_cells_compared']:,}\n")
            f.write(f"  Matches              : {stats['matches']:,} "
                    f"({stats['match_rate']:.2f}%)\n")
            f.write(f"  Mismatches (total)   : {stats['mismatches_total']:,}\n")
            f.write(f"    Full mismatches    : {stats['mismatches_full']:,}  "
                    f"(diff > {config.DECIMAL_MISMATCH_THRESHOLD})\n")
            f.write(f"    Decimal mismatches : {stats['mismatches_decimal']:,}  "
                    f"(diff <= {config.DECIMAL_MISMATCH_THRESHOLD})\n")
            f.write(f"  Missing (in ref)     : {stats['missing']:,}\n")
            f.write(f"  New values (extra)   : {stats['new_values']:,}\n\n")

            f.write('OUTPUT FILES:\n')
            f.write(f"  {config.REPORT_MISMATCHES_FULL}     "
                    f"({stats['mismatches_full']} rows)\n")
            f.write(f"  {config.REPORT_MISMATCHES_DECIMAL}  "
                    f"({stats['mismatches_decimal']} rows)\n")
            f.write(f"  {config.REPORT_MISSING}             "
                    f"({stats['missing']} rows)\n")
            f.write(f"  {config.REPORT_NEW_VALUES}          "
                    f"({stats['new_values']} rows)\n\n")

            f.write('ANNOTATED EXCEL COLORS:\n')
            f.write('  RED    = Full mismatch (wrong value / mapping error)\n')
            f.write('  ORANGE = Decimal mismatch (rounding difference)\n')
            f.write('  YELLOW = Missing (in reference, not in pipeline)\n')
            f.write('  GREEN  = New value (in pipeline, not in reference)\n\n')

            # --- Full mismatches ---
            if self.mismatches_full or self.mismatches_decimal:
                f.write('\u2500' * 70 + '\n')
                f.write('MISMATCHES\n')
                f.write('\u2500' * 70 + '\n')

                self._write_mismatch_detail(
                    f, 'FULL MISMATCHES', self.mismatches_full, 'RED')
                self._write_mismatch_detail(
                    f, 'DECIMAL MISMATCHES', self.mismatches_decimal, 'ORANGE')

            # --- Missing detail ---
            if self.missing:
                f.write('\u2500' * 70 + '\n')
                f.write('MISSING  (present in reference, absent in pipeline)\n')
                f.write('\u2500' * 70 + '\n')

                f.write('\n  By Section:\n')
                by_section = self.analyze_by_category(
                    self.missing, lambda r: r['Section'])
                for section, items in sorted(by_section.items()):
                    f.write(f"    {section}: {len(items)}\n")

                f.write('\n  By Quarter:\n')
                by_quarter = self.analyze_by_category(
                    self.missing, lambda r: r['Quarter'])
                for quarter, items in sorted(by_quarter.items()):
                    f.write(f"    {quarter}: {len(items)}\n")
                f.write('\n')

            # --- New values detail ---
            if self.extra:
                f.write('\u2500' * 70 + '\n')
                f.write('NEW VALUES  (present in pipeline, absent in reference)\n')
                f.write('\u2500' * 70 + '\n')

                f.write('\n  By Section:\n')
                by_section = self.analyze_by_category(
                    self.extra, lambda r: r['Section'])
                for section, items in sorted(by_section.items()):
                    f.write(f"    {section}: {len(items)}\n")
                f.write('\n')

            f.write('=' * 70 + '\n')
            f.write('Detailed CSV reports:\n')
            f.write(f"  {config.REPORT_MISMATCHES_FULL}\n")
            f.write(f"  {config.REPORT_MISMATCHES_DECIMAL}\n")
            f.write(f"  {config.REPORT_MISSING}\n")
            f.write(f"  {config.REPORT_NEW_VALUES}\n")
            f.write('=' * 70 + '\n')

    # ------------------------------------------------------------------
    # Build xlsx from loaded data (for CSV sources)
    # ------------------------------------------------------------------

    def _build_xlsx_from_data(self, codes, cells, descs, output_path):
        """
        Create an xlsx file from loaded data so it can be annotated.
        Used when the source is a CSV file.
        """
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = 'DATA'

        # Row 1: codes
        for c_idx, code in enumerate(codes):
            ws.cell(row=1, column=c_idx + 2, value=code)

        # Row 2: descriptions
        for c_idx, desc in enumerate(descs):
            if desc:
                ws.cell(row=2, column=c_idx + 2, value=desc)

        # Data rows
        quarters = sorted(set(k[0] for k in cells.keys()))
        for r_idx, quarter in enumerate(quarters):
            ws.cell(row=r_idx + 3, column=1, value=quarter)
            for c_idx in range(len(codes)):
                val = cells.get((quarter, c_idx))
                if val is not None:
                    ws.cell(row=r_idx + 3, column=c_idx + 2, value=val)

        wb.save(output_path)
        wb.close()

    # ------------------------------------------------------------------
    # Annotated Excel
    # ------------------------------------------------------------------

    def create_annotated_excel(self, source_path, output_path, label,
                               codes=None, cells=None, descs=None):
        """
        Create a copy of the DATA Excel with coloured highlights.
        Red = full mismatch, Orange = decimal mismatch,
        Yellow = missing, Green = new value.

        If source is CSV, builds an xlsx from the loaded data first.
        """
        if config.VERBOSE:
            print(f'Creating annotated {label} file: '
                  f'{os.path.basename(output_path)}')

        if source_path.lower().endswith('.xlsx'):
            import shutil
            shutil.copy2(source_path, output_path)
        elif codes is not None:
            self._build_xlsx_from_data(codes, cells, descs, output_path)
        else:
            if config.VERBOSE:
                print('  Skipped (not .xlsx and no data to build from)')
            return

        wb = openpyxl.load_workbook(output_path)
        ws = wb.active

        red_fill    = PatternFill(start_color='FFCCCC', end_color='FFCCCC',
                                  fill_type='solid')
        orange_fill = PatternFill(start_color='FFD699', end_color='FFD699',
                                  fill_type='solid')
        yellow_fill = PatternFill(start_color='FFFFCC', end_color='FFFFCC',
                                  fill_type='solid')
        green_fill  = PatternFill(start_color='CCFFCC', end_color='CCFFCC',
                                  fill_type='solid')
        bold_font   = Font(bold=True)

        # Map quarter labels to row numbers (Row 3+ in SIFA format)
        quarter_to_row = {}
        for r in range(3, ws.max_row + 1):
            val = ws.cell(row=r, column=1).value
            if val:
                quarter_to_row[str(val).strip()] = r

        highlight_counts = {
            'FULL': 0, 'DECIMAL': 0, 'MISSING': 0, 'NEW_VALUE': 0,
        }

        # Determine which coords to use based on label
        if label == 'Pipeline':
            coord_map = {
                'FULL':      (self.full_mismatch_coords, red_fill),
                'DECIMAL':   (self.decimal_mismatch_coords, orange_fill),
                'NEW_VALUE': (self.extra_coords, green_fill),
            }
        else:  # Reference
            coord_map = {
                'FULL':    (self.full_mismatch_coords, red_fill),
                'DECIMAL': (self.decimal_mismatch_coords, orange_fill),
                'MISSING': (self.missing_coords, yellow_fill),
            }

        for status, (coords, fill) in coord_map.items():
            for quarter, col_idx in coords:
                row_num = quarter_to_row.get(quarter)
                if row_num:
                    ws.cell(row=row_num, column=col_idx + 2).fill = fill
                    highlight_counts[status] += 1

        # Insert legend row at top
        ws.insert_rows(1)
        ws.cell(row=1, column=1, value='LEGEND:').font = bold_font
        ws.cell(row=1, column=2, value='RED = Full Mismatch').fill = red_fill
        ws.cell(row=1, column=3, value='ORANGE = Decimal Mismatch').fill = orange_fill
        ws.cell(row=1, column=4, value='YELLOW = Missing').fill = yellow_fill
        ws.cell(row=1, column=5, value='GREEN = New Value').fill = green_fill
        ws.cell(row=1, column=6,
                value=f'Annotated: {datetime.now().strftime("%Y-%m-%d %H:%M:%S")}')

        wb.save(output_path)
        wb.close()

        if config.VERBOSE:
            print(f'  Highlighted: {highlight_counts["FULL"]} full (red), '
                  f'{highlight_counts["DECIMAL"]} decimal (orange), '
                  f'{highlight_counts["MISSING"]} missing (yellow), '
                  f'{highlight_counts["NEW_VALUE"]} new values (green)')

    # ------------------------------------------------------------------
    # Run
    # ------------------------------------------------------------------

    def run(self):
        print('=' * 70)
        print('SIFA COMPARISON TOOL')
        print('=' * 70)
        print()
        print(f'Pipeline Output  : {self.pipeline_path}')
        print(f'Reference File   : {self.reference_path}')
        print(f'Source Downloads  : {len(self.source_downloads)} year file(s)')
        print(f'Mismatch threshold: {config.DECIMAL_MISMATCH_THRESHOLD} '
              f'(above = full, at/below = decimal)')
        print()

        for label, path in [('Pipeline', self.pipeline_path),
                             ('Reference', self.reference_path)]:
            if not os.path.exists(path):
                print(f'ERROR: {label} file not found: {path}')
                return 1

        # Load data
        self.pipeline_codes, self.pipeline_cells, self.pipeline_descs = \
            load_data_file(self.pipeline_path, 'Pipeline Output')
        self.reference_codes, self.reference_cells, self.reference_descs = \
            load_data_file(self.reference_path, 'Reference')

        if self.pipeline_codes != self.reference_codes:
            diffs = sum(1 for a, b in zip(self.pipeline_codes, self.reference_codes)
                        if a != b)
            print(f'\nWARNING: {diffs} column codes differ between pipeline '
                  f'and reference!')

        # Compare
        self.perform_comparison()

        # Verify mismatches against source Excel
        self.verify_mismatches()

        # Build run output directory
        timestamp   = datetime.now().strftime('%Y%m%d_%H%M%S')
        compare_dir = os.path.dirname(os.path.abspath(__file__))
        run_dir     = os.path.join(compare_dir, f'run_{timestamp}')

        os.makedirs(run_dir, exist_ok=True)

        if config.VERBOSE:
            print(f'\nOutput directory: {run_dir}')

        # ---------- Annotated Excel ----------
        if config.VERBOSE:
            print()

        # Pipeline annotated (build xlsx from CSV data if needed)
        pipeline_base = os.path.splitext(os.path.basename(self.pipeline_path))[0]
        annotated_pipeline = os.path.join(
            run_dir, f'ANNOTATED_Pipeline_{pipeline_base}.xlsx')
        self.create_annotated_excel(
            self.pipeline_path, annotated_pipeline, 'Pipeline',
            codes=self.pipeline_codes, cells=self.pipeline_cells,
            descs=self.pipeline_descs)

        # Reference annotated (build xlsx from CSV data if needed)
        reference_base = os.path.splitext(os.path.basename(self.reference_path))[0]
        annotated_ref = os.path.join(
            run_dir, f'ANNOTATED_Reference_{reference_base}.xlsx')
        self.create_annotated_excel(
            self.reference_path, annotated_ref, 'Reference',
            codes=self.reference_codes, cells=self.reference_cells,
            descs=self.reference_descs)

        # ---------- Separate CSV reports ----------
        if config.VERBOSE:
            print('\nWriting separate CSV reports...')

        self._write_category_csv(
            self.mismatches_full,
            os.path.join(run_dir, config.REPORT_MISMATCHES_FULL),
            'Full Mismatches',
        )
        self._write_category_csv(
            self.mismatches_decimal,
            os.path.join(run_dir, config.REPORT_MISMATCHES_DECIMAL),
            'Decimal Mismatches',
        )
        self._write_category_csv(
            self.missing,
            os.path.join(run_dir, config.REPORT_MISSING),
            'Missing',
        )
        self._write_category_csv(
            self.extra,
            os.path.join(run_dir, config.REPORT_NEW_VALUES),
            'New Values',
        )

        # ---------- Summary ----------
        summary_path = os.path.join(run_dir, config.SUMMARY_FILENAME)
        self.write_summary_report(summary_path)

        # ---------- Console summary ----------
        stats = self.generate_summary_stats()
        print('\n' + '=' * 70)
        print('COMPARISON COMPLETE')
        print('=' * 70)
        print(f"\nMatches              : {stats['matches']:,} ({stats['match_rate']:.2f}%)")
        print(f"Full mismatches      : {stats['mismatches_full']:,}  (RED)")
        print(f"Decimal mismatches   : {stats['mismatches_decimal']:,}  (ORANGE)")
        print(f"Missing              : {stats['missing']:,}  (YELLOW)")
        print(f"New values           : {stats['new_values']:,}  (GREEN)")
        print(f'\nOutput directory: {run_dir}')

        print('\nFiles created:')
        created = [
            f'ANNOTATED_Pipeline_{pipeline_base}.xlsx',
            f'ANNOTATED_Reference_{reference_base}.xlsx',
        ]
        created += [
            config.REPORT_MISMATCHES_FULL,
            config.REPORT_MISMATCHES_DECIMAL,
            config.REPORT_MISSING,
            config.REPORT_NEW_VALUES,
            config.SUMMARY_FILENAME,
        ]
        for name in created:
            print(f'  {name}')

        print('=' * 70)
        return 0


# ---------------------------------------------------------------------------
# Entry point
# ---------------------------------------------------------------------------

def main():
    print('=' * 60)
    print(' SIFA Compare  --  drop reference file and run me')
    print('=' * 60)
    print()

    try:
        compare_dir  = os.path.dirname(os.path.abspath(__file__))
        project_root = os.path.dirname(compare_dir)

        print('Discovering paths...')

        pipeline_file    = find_pipeline_output(project_root)
        source_downloads = find_source_downloads(project_root)
        reference_file   = find_reference_file(compare_dir)

        print(f'  Pipeline:   {os.path.basename(pipeline_file)}')
        print(f'  Reference:  {os.path.basename(reference_file)}')
        print(f'  Source DLs: {len(source_downloads)} year file(s)')

        report = ComparisonReport(
            pipeline_file, reference_file, source_downloads,
        )
        return report.run()

    except (FileNotFoundError, ValueError) as e:
        print(f'\nERROR: {e}')
        return 1

    except KeyboardInterrupt:
        print('\n\nInterrupted by user')
        return 130

    except Exception as e:
        print(f'\nERROR: {e}')
        import traceback
        traceback.print_exc()
        return 1


if __name__ == '__main__':
    sys.exit(main())
