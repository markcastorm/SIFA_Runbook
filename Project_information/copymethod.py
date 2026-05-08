import openpyxl
import os
import shutil
from datetime import datetime

# --- CONFIGURATION ---
INPUT_FOLDER = "Input"
OUTPUT_BASE_FOLDER = "Output"
SOURCE_SHEET_NAME = "2026"  # As seen in your screenshots
TARGET_SHEET_NAME = "2026_Values_Only"

def process_excel_files():
    # 1. Ensure Input folder exists
    if not os.path.exists(INPUT_FOLDER):
        os.makedirs(INPUT_FOLDER)
        print(f"Created {INPUT_FOLDER} folder. Please drop your files there and run again.")
        return

    # 2. Find all .xlsx files in the Input folder
    files_to_process = [f for f in os.listdir(INPUT_FOLDER) if f.endswith('.xlsx') and not f.startswith('~$')]

    if not files_to_process:
        print("No Excel files found in the Input folder.")
        return

    # 3. Create a unique timestamped subfolder for this specific run
    timestamp = datetime.now().strftime("%Y-%m-%d_%H-%M-%S")
    run_output_dir = os.path.join(OUTPUT_BASE_FOLDER, f"Run_{timestamp}")
    os.makedirs(run_output_dir, exist_ok=True)
    
    print(f"Starting run: {timestamp}")
    print(f"Output will be saved to: {run_output_dir}")
    print("-" * 30)

    for filename in files_to_process:
        input_path = os.path.join(INPUT_FOLDER, filename)
        output_path = os.path.join(run_output_dir, filename)
        
        print(f"Processing: {filename}...")

        try:
            # STEP A: Load with data_only=True to capture the formula results (the numbers)
            wb_data = openpyxl.load_workbook(input_path, data_only=True)
            
            if SOURCE_SHEET_NAME not in wb_data.sheetnames:
                print(f"  [Skip] Sheet '{SOURCE_SHEET_NAME}' not found in {filename}")
                continue

            source_ws = wb_data[SOURCE_SHEET_NAME]

            # STEP B: Create a new sheet and copy values
            # We create the sheet in the same workbook object
            target_ws = wb_data.create_sheet(title=TARGET_SHEET_NAME)

            # Iterate through all cells in the source sheet
            for row in source_ws.iter_rows():
                for cell in row:
                    # Copy value to the same location in the new sheet
                    target_ws.cell(row=cell.row, column=cell.column, value=cell.value)

            # STEP C: Save the result into the timestamped output folder
            wb_data.save(output_path)
            print(f"  [Success] Saved to {output_path}")

        except Exception as e:
            print(f"  [Error] Could not process {filename}: {e}")

    print("-" * 30)
    print("Process Complete.")

if __name__ == "__main__":
    process_excel_files()